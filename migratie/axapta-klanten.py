"""Zet de Axapta-export CustTable (klanten) om naar migratie/klanten-import.json.

In de export zijn de regeleinden uit het adres verdwenen ("Loswal 54271 BA DUSSEN").
Waar Axapta straat/postcode/plaats los heeft, gebruiken we die; anders splitsen we
het adres op de postcode. Lukt dat niet, dan krijgt de klant `adresControleren: true`
en staat het oorspronkelijke adres in `adresAxapta`.

Staat er een tabblad "Alternatieve adressen" (Axapta-tabel Address) in het bestand, dan
komen die adressen als `adressen` bij de klant. Koppeling in Axapta: AddrTableId 77 is
CustTable en AddrRecId is de RecId van de klant. Adressen van verwijderde klanten
(RecId niet meer in de klantentabel) en van andere tabellen (505 = leveranciers) slaan we over.

Gebruik:  python migratie/axapta-klanten.py "<Klantentabel.xlsx>"
"""
import json
import re
import sys
from collections import Counter

import openpyxl

LANDNAMEN = {
    'nederland': 'Nederland', 'the netherlands': 'Nederland', 'netherlands': 'Nederland',
    'belgie': 'België', 'belgië': 'België', 'belgium': 'België', 'belgique': 'België',
    'duitsland': 'Duitsland', 'deutschland': 'Duitsland', 'germany': 'Duitsland',
    'spain': 'Spanje', 'spanje': 'Spanje', 'espana': 'Spanje',
    'czech republic': 'Tsjechië', 'tsjechie': 'Tsjechië',
    'france': 'Frankrijk', 'frankrijk': 'Frankrijk',
    'united kingdom': 'Verenigd Koninkrijk', 'uk': 'Verenigd Koninkrijk', 'england': 'Verenigd Koninkrijk',
    'sweden': 'Zweden', 'zweden': 'Zweden', 'poland': 'Polen', 'polen': 'Polen',
    'italy': 'Italië', 'italia': 'Italië', 'austria': 'Oostenrijk', 'switzerland': 'Zwitserland',
    'denmark': 'Denemarken', 'norway': 'Noorwegen', 'noorwegen': 'Noorwegen',
    'roemenië': 'Roemenië', 'roemenie': 'Roemenië', 'romania': 'Roemenië',
}

# Postcodevorm per land, voor adressen zonder landcode voor de postcode.
POSTCODE_PER_LAND = {
    'België': r'\d{4}', 'Zwitserland': r'\d{4}', 'Oostenrijk': r'\d{4}', 'Noorwegen': r'\d{4}',
    'Denemarken': r'\d{4}', 'Duitsland': r'\d{5}', 'Spanje': r'\d{5}', 'Frankrijk': r'\d{5}',
    'Tsjechië': r'\d{3}\s?\d{2}', 'Polen': r'\d{2}-\d{3}', 'Roemenië': r'\d{6}',
}
LANDCODES = {
    'NL': 'Nederland', 'DE': 'Duitsland', 'D': 'Duitsland', 'BE': 'België', 'B': 'België',
    'E': 'Spanje', 'ES': 'Spanje', 'F': 'Frankrijk', 'FR': 'Frankrijk', 'GB': 'Verenigd Koninkrijk',
    'UK': 'Verenigd Koninkrijk', 'SE': 'Zweden', 'S': 'Zweden', 'CZ': 'Tsjechië', 'PL': 'Polen',
    'PO': 'Polen', 'A': 'Oostenrijk', 'AT': 'Oostenrijk', 'CH': 'Zwitserland', 'I': 'Italië',
    'IT': 'Italië', 'DK': 'Denemarken',
}
TAAL = {'nl': 'nl', 'de': 'de', 'en-gb': 'en', 'en-us': 'en', 'en': 'en'}


def tekst(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def haal_land_eraf(s):
    """'GenkBelgie' -> ('Genk', 'België'); ook met spatie ervoor."""
    laag = s.lower()
    for naam in sorted(LANDNAMEN, key=len, reverse=True):
        if laag.endswith(naam) and len(s) > len(naam):
            return s[: -len(naam)].rstrip(' ,'), LANDNAMEN[naam]
    return s, ''


def splits_adres(adres):
    """Geeft (straat, postcode, plaats, land, gelukt)."""
    s = re.sub(r'\s+T:\s*[\d\s-]+$', '', adres.strip())  # "… T: 013-5131284"
    s, land = haal_land_eraf(s)

    # Buitenland met landcode: "In Der Illekatte 7-9D-58339 BRECKERFELD", "Taunusweg 6B-3600 Genk",
    # "Schiesstrasse 68D-40549Dusseldorf", "Heihoefke 4bB-2960  Brecht"
    m = list(re.finditer(r'(?<![A-Z])([A-Z]{1,2})-\s?(\d{4,5})\s*([A-Za-zÀ-ž].*)$', s))
    if m:
        m = m[-1]
        straat = s[: m.start()].strip()
        if straat:
            return straat, f'{m.group(1)}-{m.group(2)}', m.group(3).strip(), \
                land or LANDCODES.get(m.group(1), ''), True

    # Nederland: "Postbus 1257050 AC VARSSEVELD", "Gooiland 21948RC Beverwijk", "853125 BA, Schiedam"
    if land in ('', 'Nederland'):
        m = list(re.finditer(r'(\d{4})\s{0,2}([A-Z]{2}),?\s+(\S.*)$', s))
        if m:
            m = m[-1]
            straat = s[: m.start()].strip()
            if straat:
                return straat, f'{m.group(1)} {m.group(2)}', m.group(3).strip(), 'Nederland', True

    # Buitenland zonder landcode, als het land bekend is: "Vennekeslaan 13665 As" (België)
    vorm = POSTCODE_PER_LAND.get(land)
    if vorm:
        m = list(re.finditer(rf'({vorm})\s*([A-Za-zÀ-ž].*)$', s))
        if m:
            m = m[-1]
            straat = s[: m.start()].strip()
            if straat:
                return straat, m.group(1), m.group(2).strip(), land, True

    return s, '', '', land, False


# Adressoort (veld "type" in Axapta) -> `type` in fmid-erp, zie ADRES_TYPES in src/lib/stamgegevens.js.
ADRES_TYPES = {'Levering': 'levering', 'Factuur': 'factuur', 'Alternatief afleveradres': 'alternatief'}
TABEL_KLANTEN = 77


def lees_tabblad(ws):
    """Geeft (rijen, v) met v(rij, veldnaam) -> tekst; bij dubbele kolomnamen telt de eerste."""
    rijen = ws.iter_rows(values_only=True)
    ix = {}
    for i, k in enumerate(next(rijen)):
        if k and k not in ix:
            ix[k] = i

    def v(r, veld):
        return tekst(r[ix[veld]]) if veld in ix else ''

    return rijen, v


def adres_uit(v, r, standaard_land=''):
    """Straat/postcode/plaats/land uit een Axapta-rij met Address, Street, ZipCode, City, Country."""
    adres = v(r, 'Address')
    if v(r, 'ZipCode') and v(r, 'City'):
        straat, postcode, plaats = v(r, 'Street'), v(r, 'ZipCode'), v(r, 'City')
        land = LANDCODES.get(v(r, 'Country').upper(), '') or haal_land_eraf(adres)[1]
        gelukt = True
    else:
        straat, postcode, plaats, land, gelukt = splits_adres(adres)
    if not land:
        land = LANDCODES.get(v(r, 'Country').upper(), '') or standaard_land
    return adres, straat, postcode, plaats, land, gelukt


def lees_alternatieve_adressen(ws):
    """Geeft {RecId klant: [adres, ...]}."""
    rijen, v = lees_tabblad(ws)
    per_klant = {}
    for r in rijen:
        if v(r, 'AddrTableId') != str(TABEL_KLANTEN):
            continue
        if not (v(r, 'Address') or v(r, 'Street')):
            continue  # alleen een naam en geen adres: maakt Axapta vaak zelf aan bij factuur
        adres, straat, postcode, plaats, land, gelukt = adres_uit(v, r)
        per_klant.setdefault(v(r, 'AddrRecId'), []).append({
            'type': ADRES_TYPES.get(v(r, 'type'), 'overig'),
            'naam': v(r, 'Name'),
            'straat': straat,
            'postcode': postcode,
            'plaats': plaats,
            'land': land,
            'telefoon': v(r, 'Phone'),
            'email': v(r, 'Email'),
            'adresAxapta': adres,
            'adresControleren': not gelukt,
        })
    return per_klant


def main(pad):
    wb = openpyxl.load_workbook(pad, read_only=True, data_only=True)
    ws = wb['Klantentabel'] if 'Klantentabel' in wb.sheetnames else wb.active
    rijen, v = lees_tabblad(ws)
    alt = lees_alternatieve_adressen(wb['Alternatieve adressen']) if 'Alternatieve adressen' in wb.sheetnames else {}

    uit = []
    for r in rijen:
        code = v(r, 'AccountNum')
        if not code:
            continue
        adres, straat, postcode, plaats, land, gelukt = adres_uit(
            v, r, 'Nederland' if v(r, 'CustGroup') == 'NL' else '')

        btw_nummer = v(r, 'VATNum')
        if '?' in btw_nummer or '*' in btw_nummer:
            btw_nummer = ''
        btw_groep = v(r, 'TaxGroup') or {'NL': 'NL', 'EU': 'EU-IC' if btw_nummer else 'EU-belast',
                                          'Niet EU': 'Non-EU'}.get(v(r, 'CustGroup'), 'NL')
        termijn = re.match(r'\d+', v(r, 'PaymTermId'))

        uit.append({
            'klantcode': code,
            'naam': v(r, 'Name'),
            'zoeknaam': v(r, 'NameAlias'),
            'klantgroep': v(r, 'CustGroup'),
            'contactpersoon': '',
            'telefoon': v(r, 'Phone'),
            'mobiel': v(r, 'CellularPhone'),
            'fax': v(r, 'TeleFax'),
            'email': v(r, 'Email'),
            'website': v(r, 'URL'),
            'straat': straat,
            'postcode': postcode,
            'plaats': plaats,
            'land': land,
            'adresAxapta': adres,
            'adresControleren': not gelukt,
            'btwGroep': btw_groep,
            'btwNummer': btw_nummer,
            'inclBtw': v(r, 'InclTax') == 'Ja',
            'valuta': v(r, 'Currency') or 'EUR',
            'betalingstermijn': int(termijn.group()) if termijn else (0 if v(r, 'PaymTermId') == 'Kontant' else 30),
            'betalingsconditie': v(r, 'PaymTermId'),
            'leveringsvoorwaarde': v(r, 'DlvTerm'),
            'leveringswijze': v(r, 'DlvMode'),
            'taal': TAAL.get(v(r, 'LanguageId').lower(), 'nl'),
            'kvkNummer': '',
            'geblokkeerd': v(r, 'Blocked') == 'Ja',
            # Land onbekend bij een alternatief adres: dan meestal hetzelfde land als de klant.
            'adressen': [{**a, 'land': a['land'] or land} for a in alt.get(v(r, 'RecId'), [])],
        })

    with open('migratie/klanten-import.json', 'w', encoding='utf-8') as f:
        json.dump(uit, f, ensure_ascii=False, indent=1)

    print('klanten:', len(uit))
    print('adres na te kijken:', sum(k['adresControleren'] for k in uit))
    print('landen:', Counter(k['land'] for k in uit))
    print('btw-groepen:', Counter(k['btwGroep'] for k in uit))
    print('talen:', Counter(k['taal'] for k in uit))
    adressen = [a for k in uit for a in k['adressen']]
    print('alternatieve adressen:', len(adressen), Counter(a['type'] for a in adressen))
    print('  bij klanten:', sum(1 for k in uit if k['adressen']),
          '· na te kijken:', sum(a['adresControleren'] for a in adressen))
    for k in uit:
        if k['adresControleren']:
            print('  NAKIJKEN', k['klantcode'], k['naam'], '|', k['adresAxapta'])


if __name__ == '__main__':
    main(sys.argv[1])
